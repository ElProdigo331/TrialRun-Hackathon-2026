"""
Generate Brain_Oil.pptx PowerPoint presentation for Energy AI Hackathon 2026.
Following the official 6-slide hackathon template format.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def add_title_slide(prs, title, subtitle, team_members):
    """Slide 1: Title slide with team info."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(40)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER
    
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(9), Inches(0.5))
    sub_frame = sub_box.text_frame
    sub_para = sub_frame.paragraphs[0]
    sub_para.text = subtitle
    sub_para.font.size = Pt(20)
    sub_para.alignment = PP_ALIGN.CENTER
    
    team_label = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(0.4))
    team_label_frame = team_label.text_frame
    team_label_para = team_label_frame.paragraphs[0]
    team_label_para.text = "Team Brain Oil"
    team_label_para.font.size = Pt(24)
    team_label_para.font.bold = True
    team_label_para.alignment = PP_ALIGN.CENTER
    
    team_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.0), Inches(9), Inches(1))
    team_frame = team_box.text_frame
    team_para = team_frame.paragraphs[0]
    team_para.text = team_members
    team_para.font.size = Pt(16)
    team_para.alignment = PP_ALIGN.CENTER
    
    uni_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(9), Inches(0.5))
    uni_frame = uni_box.text_frame
    uni_para = uni_frame.paragraphs[0]
    uni_para.text = "Hildebrand Department of Petroleum and Geosystems Engineering"
    uni_para.font.size = Pt(14)
    uni_para.alignment = PP_ALIGN.CENTER
    
    date_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.4), Inches(9), Inches(0.5))
    date_frame = date_box.text_frame
    date_para = date_frame.paragraphs[0]
    date_para.text = "The University of Texas at Austin"
    date_para.font.size = Pt(14)
    date_para.alignment = PP_ALIGN.CENTER

def add_executive_summary_slide(prs):
    """Slide 2: Executive Summary answering the 4 key questions."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Executive Summary"
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    content = [
        ("1. The Problem", "Predict 3-year cumulative oil production for 12 preproduction wells (IDs 72-83) with uncertainty quantification (100 realizations)."),
        ("2. Our Solution", "Ridge Regression (alpha=1.0) with StandardScaler, stepwise feature selection (105→10), and Bagging Ensemble for uncertainty."),
        ("3. What We Learned", "For small datasets (n=71), simple regularized models beat complex ones. Porosity (φ) is the primary driver of production."),
        ("4. Our Results", "Test R² = 0.9905 (EXCELLENT), RMSE = 1.57M BBL (4.7% error). Ridge outperformed Random Forest and XGBoost.")
    ]
    
    y_pos = 1.0
    for label, text in content:
        label_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(9), Inches(0.3))
        label_frame = label_box.text_frame
        label_para = label_frame.paragraphs[0]
        label_para.text = label
        label_para.font.size = Pt(16)
        label_para.font.bold = True
        label_para.font.color.rgb = RGBColor(0, 102, 153)
        
        text_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos + 0.3), Inches(9), Inches(0.8))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        text_para = text_frame.paragraphs[0]
        text_para.text = text
        text_para.font.size = Pt(14)
        
        y_pos += 1.2

def add_workflow_overview_slide(prs):
    """Slide 3: Workflow Overview with steps."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Workflow Overview"
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    steps = [
        "1. Data Loading → Load well logs (71 train, 12 test) + sand proportion map",
        "2. MICE + CART Imputation → Fill missing values at depth level (Van Buuren 2018)",
        "3. Aggregation → Multi-row depth data → One row per well (mean, std, min, max)",
        "4. Feature Engineering → 105 features including RQI, FZI, spatial features",
        "5. Feature Selection → Correlation filter (105→61) + Stepwise (61→10)",
        "6. Model Training → Ridge Regression with StandardScaler (alpha=1.0)",
        "7. Uncertainty → Bagging Ensemble (100 estimators) for R1-R100 realizations"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.0), Inches(9.2), Inches(5.5))
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    
    for i, step in enumerate(steps):
        if i == 0:
            p = content_frame.paragraphs[0]
        else:
            p = content_frame.add_paragraph()
        p.text = step
        p.font.size = Pt(16)
        p.space_before = Pt(12)
        p.space_after = Pt(4)

def add_workflow_details_slide(prs):
    """Slide 4: Workflow Details - Key Decisions and Innovations."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Key Modeling Decisions"
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    left_title = slide.shapes.add_textbox(Inches(0.3), Inches(1.0), Inches(4.5), Inches(0.4))
    left_title_frame = left_title.text_frame
    left_title_para = left_title_frame.paragraphs[0]
    left_title_para.text = "Why Ridge Regression?"
    left_title_para.font.size = Pt(18)
    left_title_para.font.bold = True
    left_title_para.font.color.rgb = RGBColor(0, 102, 153)
    
    left_content = [
        "• n=71 is small → regularized linear beats trees",
        "• Hastie et al. (2009): lower variance for small n",
        "• Ridge R²=0.99 vs RF R²=0.85 vs XGB R²=0.82",
        "• Alpha=1.0 provides balanced regularization"
    ]
    
    left_box = slide.shapes.add_textbox(Inches(0.3), Inches(1.4), Inches(4.5), Inches(2))
    left_frame = left_box.text_frame
    left_frame.word_wrap = True
    for i, point in enumerate(left_content):
        if i == 0:
            p = left_frame.paragraphs[0]
        else:
            p = left_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(14)
        p.space_before = Pt(6)
    
    right_title = slide.shapes.add_textbox(Inches(5.2), Inches(1.0), Inches(4.5), Inches(0.4))
    right_title_frame = right_title.text_frame
    right_title_para = right_title_frame.paragraphs[0]
    right_title_para.text = "Why Porosity (φ) Matters Most?"
    right_title_para.font.size = Pt(18)
    right_title_para.font.bold = True
    right_title_para.font.color.rgb = RGBColor(0, 102, 153)
    
    right_content = [
        "• φ directly measures storage capacity",
        "• OOIP = 7758 × A × h × φ × (1-Sw) / Bo",
        "• Higher porosity = more oil = more production",
        "• Domain experts prioritize φ over k"
    ]
    
    right_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.4), Inches(4.5), Inches(2))
    right_frame = right_box.text_frame
    right_frame.word_wrap = True
    for i, point in enumerate(right_content):
        if i == 0:
            p = right_frame.paragraphs[0]
        else:
            p = right_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(14)
        p.space_before = Pt(6)
    
    config_title = slide.shapes.add_textbox(Inches(0.3), Inches(3.6), Inches(9), Inches(0.4))
    config_title_frame = config_title.text_frame
    config_title_para = config_title_frame.paragraphs[0]
    config_title_para.text = "Final Configuration (57+ configs tested)"
    config_title_para.font.size = Pt(18)
    config_title_para.font.bold = True
    config_title_para.font.color.rgb = RGBColor(0, 102, 153)
    
    config_text = "Model: Ridge (α=1.0) | Normalize: StandardScaler | Features: 10 (stepwise) | Sand Map: Smooth 3x3 | Uncertainty: Bagging (100)"
    config_box = slide.shapes.add_textbox(Inches(0.3), Inches(4.0), Inches(9.4), Inches(0.6))
    config_frame = config_box.text_frame
    config_frame.word_wrap = True
    config_para = config_frame.paragraphs[0]
    config_para.text = config_text
    config_para.font.size = Pt(14)

def add_results_slide(prs):
    """Slide 5: Results and Discussions."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Results and Discussions"
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    perf_title = slide.shapes.add_textbox(Inches(0.3), Inches(1.0), Inches(4.5), Inches(0.4))
    perf_title_frame = perf_title.text_frame
    perf_title_para = perf_title_frame.paragraphs[0]
    perf_title_para.text = "Model Performance"
    perf_title_para.font.size = Pt(18)
    perf_title_para.font.bold = True
    perf_title_para.font.color.rgb = RGBColor(0, 102, 153)
    
    perf_content = [
        "• Test R² = 0.9905 (EXCELLENT ≥0.93)",
        "• CV R² = 0.9539 ± 0.0456 (stable)",
        "• RMSE = 1.57M BBL (4.7% of mean)",
        "• All metrics exceed industry benchmarks"
    ]
    
    perf_box = slide.shapes.add_textbox(Inches(0.3), Inches(1.4), Inches(4.5), Inches(2))
    perf_frame = perf_box.text_frame
    perf_frame.word_wrap = True
    for i, point in enumerate(perf_content):
        if i == 0:
            p = perf_frame.paragraphs[0]
        else:
            p = perf_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(14)
        p.space_before = Pt(6)
    
    pred_title = slide.shapes.add_textbox(Inches(5.2), Inches(1.0), Inches(4.5), Inches(0.4))
    pred_title_frame = pred_title.text_frame
    pred_title_para = pred_title_frame.paragraphs[0]
    pred_title_para.text = "Predictions (Wells 72-83)"
    pred_title_para.font.size = Pt(18)
    pred_title_para.font.bold = True
    pred_title_para.font.color.rgb = RGBColor(0, 102, 153)
    
    pred_content = [
        "• 12 point estimates + 100 realizations each",
        "• Range: ~15M to ~47M BBL over 3 years",
        "• Uncertainty via Bagging Ensemble",
        "• Output: solution.csv (102 columns)"
    ]
    
    pred_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.4), Inches(4.5), Inches(2))
    pred_frame = pred_box.text_frame
    pred_frame.word_wrap = True
    for i, point in enumerate(pred_content):
        if i == 0:
            p = pred_frame.paragraphs[0]
        else:
            p = pred_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(14)
        p.space_before = Pt(6)
    
    comp_title = slide.shapes.add_textbox(Inches(0.3), Inches(3.5), Inches(9), Inches(0.4))
    comp_title_frame = comp_title.text_frame
    comp_title_para = comp_title_frame.paragraphs[0]
    comp_title_para.text = "Model Comparison (57+ configurations tested)"
    comp_title_para.font.size = Pt(18)
    comp_title_para.font.bold = True
    comp_title_para.font.color.rgb = RGBColor(0, 102, 153)
    
    comp_text = "Ridge R²=0.9905 (WINNER) | Random Forest R²=0.85 | XGBoost R²=0.82 | Linear Regression: Overfit"
    comp_box = slide.shapes.add_textbox(Inches(0.3), Inches(3.9), Inches(9.4), Inches(0.6))
    comp_frame = comp_box.text_frame
    comp_frame.word_wrap = True
    comp_para = comp_frame.paragraphs[0]
    comp_para.text = comp_text
    comp_para.font.size = Pt(14)

def add_feedback_slide(prs):
    """Slide 6: Feedback - What learned, liked, improvements."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Feedback"
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    sections = [
        ("What We Learned", [
            "• Simple models win for small datasets (n<100)",
            "• Domain knowledge (porosity priority) improves feature selection",
            "• MICE+CART imputation preserves data relationships",
            "• Bagging provides robust uncertainty quantification"
        ]),
        ("What We Liked", [
            "• Real-world petroleum engineering problem",
            "• Comprehensive workshop materials (Prof. Pyrcz & Prof. Foster)",
            "• Multi-disciplinary team collaboration opportunity"
        ]),
        ("Suggestions for Next Year", [
            "• More time for model validation and iteration",
            "• Additional spatial/temporal data for enhanced predictions"
        ])
    ]
    
    y_pos = 1.0
    for section_title, points in sections:
        sec_title = slide.shapes.add_textbox(Inches(0.3), Inches(y_pos), Inches(9), Inches(0.4))
        sec_title_frame = sec_title.text_frame
        sec_title_para = sec_title_frame.paragraphs[0]
        sec_title_para.text = section_title
        sec_title_para.font.size = Pt(16)
        sec_title_para.font.bold = True
        sec_title_para.font.color.rgb = RGBColor(0, 102, 153)
        
        content_box = slide.shapes.add_textbox(Inches(0.3), Inches(y_pos + 0.35), Inches(9.4), Inches(1.5))
        content_frame = content_box.text_frame
        content_frame.word_wrap = True
        for i, point in enumerate(points):
            if i == 0:
                p = content_frame.paragraphs[0]
            else:
                p = content_frame.add_paragraph()
            p.text = point
            p.font.size = Pt(13)
            p.space_before = Pt(3)
        
        y_pos += 0.35 + len(points) * 0.35 + 0.3

def create_presentation():
    """Create the complete 6-slide presentation."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    team_members = "Kailasadatta Boggaram, Jayanth Damodaran, Bilal Shihab, Carlos Fabela"
    
    add_title_slide(prs, 
                   "Oil Production Prediction", 
                   "Energy AI Hackathon 2026 - January 25th, 2026",
                   team_members)
    
    add_executive_summary_slide(prs)
    add_workflow_overview_slide(prs)
    add_workflow_details_slide(prs)
    add_results_slide(prs)
    add_feedback_slide(prs)
    
    output_path = "Brain_Oil.pptx"
    prs.save(output_path)
    print(f"Created {output_path} with {len(prs.slides)} slides (matching hackathon template)")
    return output_path

if __name__ == "__main__":
    create_presentation()
